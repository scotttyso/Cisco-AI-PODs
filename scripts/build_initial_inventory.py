#!/usr/bin/env python3
"""Build an initial inventory from order spreadsheets."""

import argparse
import ipaddress
import json
import os
import re
import tempfile
from collections.abc import Iterable
from pathlib import Path
from urllib.parse import unquote, urlsplit

import yaml
from jsonschema import Draft202012Validator, FormatChecker
from openpyxl import load_workbook
from referencing import Registry, Resource
from referencing.jsonschema import DRAFT202012


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = REPOSITORY_ROOT / "host_vars/initial_inventory.ezai.yaml"
DEFAULT_SCHEMA = REPOSITORY_ROOT / "schemas/source/cisco-ai-pods.json"
HEADERS = {"pid": "pid", "serial number": "serial_number", "mac addresses": "mac"}


class IndentedDumper(yaml.SafeDumper):
    def increase_indent(self, flow=False, indentless=False):
        return super().increase_indent(flow, indentless=False)


def normalize_mac(value: str) -> str:
    mac_hex = re.sub(r"[:.-]", "", str(value).strip())
    if not re.fullmatch(r"[0-9a-fA-F]{1,12}", mac_hex):
        raise ValueError(f"Invalid MAC address: {value}")
    mac_hex = mac_hex.zfill(12)
    return ":".join(mac_hex[index:index + 2] for index in range(0, 12, 2)).lower()


def read_inventory(folder: Path) -> list[dict[str, str]]:
    return read_inventory_files(sorted(folder.glob("*.xlsx")))


def read_inventory_files(files: Iterable[Path]) -> list[dict[str, str]]:
    files = [Path(path) for path in files if not Path(path).name.startswith("~$")]
    if not files:
        raise ValueError("No .xlsx files found")

    inventory = []
    found_headers = False
    for path in files:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook:
                columns = None
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    if columns is None:
                        labels = {str(value).strip().casefold(): index for index, value in enumerate(row) if value is not None}
                        if all(header in labels for header in HEADERS):
                            columns = {key: labels[header] for header, key in HEADERS.items()}
                            found_headers = True
                        continue

                    mac = row[columns["mac"]] if len(row) > columns["mac"] else None
                    if mac is None or not str(mac).strip():
                        continue
                    try:
                        normalized_mac = normalize_mac(mac)
                    except ValueError as error:
                        raise ValueError(f"Invalid MAC address in {path.name}, {sheet.title}, row {row_number}: {mac}") from error
                    pid = row[columns["pid"]] if len(row) > columns["pid"] else None
                    serial = row[columns["serial_number"]] if len(row) > columns["serial_number"] else None
                    if not pid or not serial:
                        raise ValueError(f"Missing PID or serial number in {path.name}, {sheet.title}, row {row_number}")
                    if len(inventory) >= 254:
                        raise ValueError("More than 254 MAC entries; 198.18.0.X cannot hold them all")
                    inventory.append({
                        "hostname": "unknown",
                        "ip": f"198.18.0.{len(inventory) + 1}",
                        "mac": normalized_mac,
                        "pid": str(pid).strip(),
                        "serial_number": str(serial).strip(),
                    })
        finally:
            workbook.close()

    if not found_headers:
        raise ValueError("No sheets contain PID, Serial number, and MAC Addresses columns")
    return inventory


def schema_validator(schema_path: Path = DEFAULT_SCHEMA, definition: str = "shared_services") -> Draft202012Validator:
    schema_path = schema_path.resolve()

    def retrieve(uri: str) -> Resource:
        location = urlsplit(uri)
        if location.scheme != "file":
            raise ValueError("Schema references must be local files")
        path = Path(unquote(location.path)).resolve()
        if not path.is_relative_to(schema_path.parent):
            raise ValueError("Schema reference is outside the schema directory")
        return Resource.from_contents(json.loads(path.read_text(encoding="utf-8")), default_specification=DRAFT202012)

    schema = {"$ref": f"{schema_path.as_uri()}#/definitions/{definition}"}
    return Draft202012Validator(schema, registry=Registry(retrieve=retrieve), format_checker=FormatChecker())


def validate_shared_services(shared_services: dict, schema_path: Path = DEFAULT_SCHEMA) -> list[str]:
    errors = schema_validator(schema_path).iter_errors(shared_services)
    return [f"shared_services.{'.'.join(map(str, error.absolute_path)) or '(root)'}: {error.message}"
            for error in sorted(errors, key=lambda error: str(list(error.absolute_path)))]


def build_inventory_document(inventory: list[dict[str, str]], shared_services: dict | None = None,
                             schema_path: Path = DEFAULT_SCHEMA) -> dict:
    if not inventory:
        raise ValueError("Add at least one inventory row")
    shared_services = {} if shared_services is None else shared_services
    errors = validate_shared_services(shared_services, schema_path)
    if errors:
        raise ValueError("\n".join(errors))
    normalized = []
    macs = set()
    addresses = set()
    for position, row in enumerate(inventory, start=1):
        entry = {key: str(row.get(key, "")).strip() for key in ("hostname", "ip", "mac", "pid", "serial_number")}
        if any(not value for value in entry.values()):
            raise ValueError(f"Row {position}: all inventory fields are required")
        if len(entry["hostname"]) > 253 or not all(
            re.fullmatch(r"[A-Za-z0-9](?:[A-Za-z0-9-]{0,61}[A-Za-z0-9])?", label)
            for label in entry["hostname"].split(".")):
            raise ValueError(f"Row {position}: invalid hostname")
        try:
            entry["ip"] = str(ipaddress.ip_address(entry["ip"]))
            entry["mac"] = normalize_mac(entry["mac"])
        except ValueError as error:
            raise ValueError(f"Row {position}: invalid IP or MAC address") from error
        if entry["mac"] in macs:
            raise ValueError(f"Row {position}: duplicate MAC address")
        if entry["ip"] in addresses:
            raise ValueError(f"Row {position}: duplicate IP address")
        macs.add(entry["mac"])
        addresses.add(entry["ip"])
        normalized.append(entry)
    return {"initial_inventory": normalized, "shared_services": shared_services}


def write_inventory(inventory: list[dict[str, str]], shared_services: dict, output: Path = DEFAULT_OUTPUT,
                    schema_path: Path = DEFAULT_SCHEMA) -> dict:
    document = build_inventory_document(inventory, shared_services, schema_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(prefix=f".{output.name}.", dir=output.parent)
    temporary = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as stream:
            yaml.dump(document, stream, Dumper=IndentedDumper, sort_keys=False, explicit_start=True)
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)
    return document


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="YAML output path")
    parser.add_argument("--folder", type=Path, help="Folder containing order workbooks")
    parser.add_argument("--inventory", type=Path, help="Edited inventory YAML containing initial_inventory")
    parser.add_argument("--shared-services", type=Path, help="Shared services YAML, optionally wrapped in shared_services")
    parser.add_argument("--schema", type=Path, default=DEFAULT_SCHEMA, help="Source schema path")
    arguments = parser.parse_args()

    try:
        if arguments.inventory:
            inventory = yaml.safe_load(arguments.inventory.read_text(encoding="utf-8"))["initial_inventory"]
        else:
            folder = arguments.folder
            if folder is None:
                default_folder = Path.home() / "orders"
                answer = input(f"Folder containing .xlsx files [{default_folder}]: ").strip()
                folder = Path(answer).expanduser() if answer else default_folder
            if not folder.is_dir():
                raise ValueError(f"Not a directory: {folder}")
            inventory = read_inventory(folder)
        shared_services = {}
        if arguments.shared_services:
            settings = yaml.safe_load(arguments.shared_services.read_text(encoding="utf-8")) or {}
            if not isinstance(settings, dict):
                raise ValueError("Shared services YAML must contain a mapping")
            shared_services = settings.get("shared_services", settings)
        write_inventory(inventory, shared_services, arguments.output, arguments.schema)
    except (OSError, ValueError, KeyError, TypeError, yaml.YAMLError) as error:
        parser.error(str(error))

    print(f"Wrote {len(inventory)} entries to {arguments.output}")


if __name__ == "__main__":
    main()